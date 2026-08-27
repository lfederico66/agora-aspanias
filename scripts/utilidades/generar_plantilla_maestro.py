"""
generar_plantilla_maestro.py — Genera la plantilla maestra de captura de datos
de personas atendidas con los campos REALES del modelo Django de ÁGORA.

Uso:
    python scripts/utilidades/generar_plantilla_maestro.py <salida.xlsx>

7 hojas espejo de las entidades: Persona (1 fila/persona) + hojas de detalle
(varias filas/persona, enlazadas por DNI o código interno). Sin datos reales:
solo una fila de EJEMPLO FICTICIO por hoja, en gris.

Al evolucionar los modelos (backend/personas/models/), regenerar y redistribuir.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

VERDE = PatternFill("solid", fgColor="047857")
AMBAR = PatternFill("solid", fgColor="FDE68A")
GRIS = Font(color="9CA3AF", italic=True)
BLANCO = Font(color="FFFFFF", bold=True)

# (título hoja, [(campo, obligatorio, desplegable|None)], [fila ejemplo])
HOJAS = [
    ("1 Persona", [
        ("codigo_interno *", True, None),
        ("nombre *", True, None),
        ("apellido_1 *", True, None),
        ("apellido_2", False, None),
        ("dni_nie", False, None),
        ("fecha_nacimiento * (dd/mm/aaaa)", True, None),
        ("sexo", False, '"M,F,X,No informa"'),
        ("nacionalidad (ISO-2)", False, None),
        ("direccion_calle", False, None),
        ("direccion_cp", False, None),
        ("direccion_municipio", False, None),
        ("direccion_provincia", False, None),
        ("nuss (nº Seguridad Social)", False, None),
        ("tsi (tarjeta sanitaria)", False, None),
        ("tsi_caducidad", False, None),
        ("centro_salud", False, None),
        ("forma_comunicacion_preferente", False, '"Verbal,Pictogramas,Lengua de signos,Mixta,Otra"'),
        ("idioma_preferente", False, None),
        ("usa_saac", False, '"Sí,No"'),
        ("saac_notas", False, None),
        ("corresponsable_principal", False, None),
        ("centro_referencia *", True, None),
        ("servicios_activos (separar con ;)", False, None),
        ("fecha_alta * (primera entrada al grupo)", True, None),
        ("gestor_caso", False, None),
        ("persona_referencia", False, None),
        ("tipo_discapacidad", False, '"Intelectual,Física,Sensorial,Mixta,Sin discapacidad"'),
        ("grado_discapacidad_pct", False, None),
        ("fecha_reconocimiento_discapacidad", False, None),
        ("grado_dependencia", False, '"Sin valorar,Grado I,Grado II,Grado III"'),
        ("fecha_valoracion_bvd", False, None),
        ("notas_relevantes", False, None),
    ], ["FUE-2026-00999", "Persona", "De Prueba", "Ejemplo", "00000000X", "01/01/1980",
        "F", "ES", "C/ Ficticia 1", "09001", "Burgos", "Burgos", "281234567890", "TSI-000",
        "01/01/2030", "CS Ejemplo", "Pictogramas", "es", "Sí", "Tablero de comunicación",
        "Madre", "CENTRO MAYORES FUENTECILLAS", "Residencia; Centro de Día", "15/09/2018",
        "Gestora Ejemplo", "Profesional Ejemplo", "Intelectual", "65", "10/05/2009",
        "Grado II", "12/03/2022", "EJEMPLO FICTICIO — borrar esta fila"]),
    ("2 Salud y cuidados", [
        ("dni_o_codigo *", True, None),
        ("grupo_sanguineo", False, '"Desconocido,0+,0-,A+,A-,B+,B-,AB+,AB-"'),
        ("vacunacion_permitida", False, '"Sí,No"'),
        ("motivo_no_vacunacion", False, None),
        ("contacto_medico_referencia", False, None),
        ("antecedentes_familiares", False, None),
        ("tipo_dieta", False, None),
        ("textura_alimentos", False, None),
        ("textura_liquidos", False, None),
        ("observaciones_alimentacion", False, None),
        ("cuidados: respiracion", False, None),
        ("cuidados: audicion", False, None),
        ("cuidados: vision", False, None),
        ("cuidados: alimentacion", False, None),
        ("cuidados: sueno_descanso", False, None),
        ("cuidados: eliminacion", False, None),
        ("cuidados: movilidad", False, None),
        ("cuidados: autonomia_abvd", False, None),
        ("cuidados: conducta", False, None),
        ("cuidados: cuidados_piel", False, None),
        ("cuidados: observaciones", False, None),
    ], ["00000000X", "A+", "Sí", "", "Dr. Ejemplo (CS Ejemplo)", "", "Blanda sin sal",
        "Normal", "Néctar", "", "Sin incidencias", "Audífono derecho", "Gafas",
        "Autónoma con supervisión", "Conciliación normal", "Continente diurna",
        "Bastón en exteriores", "Apoyo verbal en aseo", "Tranquila",
        "Hidratación diaria", "EJEMPLO FICTICIO — borrar"]),
    ("3 Alergias", [
        ("dni_o_codigo *", True, None),
        ("tipo *", True, '"Medicamento,Alimento,Ambiental,Otro"'),
        ("sustancia *", True, None),
        ("reaccion", False, None),
        ("gravedad", False, '"Leve,Moderada,Grave"'),
        ("fecha_diagnostico", False, None),
        ("activa", False, '"Sí,No"'),
        ("notas", False, None),
    ], ["00000000X", "Medicamento", "Penicilina", "Urticaria", "Grave", "", "Sí",
        "EJEMPLO FICTICIO — borrar"]),
    ("4 Enfermedades cronicas", [
        ("dni_o_codigo *", True, None),
        ("enfermedad *", True, None),
        ("fecha_diagnostico", False, None),
        ("especialista_referente", False, None),
        ("en_seguimiento", False, '"Sí,No"'),
        ("notas", False, None),
    ], ["00000000X", "Epilepsia focal", "14/02/2018", "Dra. Ejemplo (HUBU)", "Sí",
        "EJEMPLO FICTICIO — borrar"]),
    ("5 Medicacion", [
        ("dni_o_codigo *", True, None),
        ("nombre_comercial *", True, None),
        ("principio_activo", False, None),
        ("dosis_presentacion (p.ej. 500 mg)", False, None),
        ("via_administracion", False, '"Oral,Sublingual,Tópica,Inhalada,Intramuscular,Subcutánea,Rectal,Oftálmica,Otra"'),
        ("dosis_desayuno", False, None),
        ("dosis_comida", False, None),
        ("dosis_cena", False, None),
        ("dosis_acostarse", False, None),
        ("si_precisa", False, '"Sí,No"'),
        ("pauta_especial", False, None),
        ("fecha_inicio", False, None),
        ("notas", False, None),
    ], ["00000000X", "Keppra", "Levetiracetam", "500 mg", "Oral", "1", "0", "1", "",
        "No", "", "14/02/2018", "EJEMPLO FICTICIO — borrar"]),
    ("6 Contactos y familia", [
        ("dni_o_codigo persona *", True, None),
        ("nombre_contacto *", True, None),
        ("relacion *", True, '"Madre,Padre,Hermana,Hermano,Hija,Hijo,Tía,Tío,Pareja,Tutor legal,Amistad,Otro"'),
        ("telefono_fijo", False, None),
        ("movil", False, None),
        ("email", False, None),
        ("direccion", False, None),
        ("municipio", False, None),
        ("es_referencia_principal", False, '"Sí,No"'),
        ("es_emergencia", False, '"Sí,No"'),
        ("orden_emergencia (1,2,3...)", False, None),
        ("consentimiento_comunicacion", False, '"Sí,No"'),
        ("fecha_consentimiento", False, None),
        ("notificable_cerca (app familias)", False, '"Sí,No"'),
        ("notas", False, None),
    ], ["00000000X", "Contacto De Prueba", "Madre", "947000000", "600000000",
        "ejemplo@ejemplo.es", "C/ Ficticia 1", "Burgos", "Sí", "Sí", "1", "Sí",
        "01/02/2026", "Sí", "EJEMPLO FICTICIO — borrar"]),
    ("7 Medida de apoyo", [
        ("dni_o_codigo *", True, None),
        ("tipo *", True, '"Sin medida,Guarda de hecho,Curatela,Defensor judicial"'),
        ("subtipo_curatela", False, '"Asistencial,Representativa"'),
        ("fecha_resolucion", False, None),
        ("organo_judicial", False, None),
        ("numero_procedimiento", False, None),
        ("es_curatela_entidad", False, '"Sí,No"'),
        ("persona_curadora", False, None),
        ("entidad_curadora", False, None),
        ("ambito_apoyos", False, None),
        ("ambito_capacidad_conservada", False, None),
        ("vigente", False, '"Sí,No"'),
    ], ["00000000X", "Curatela", "Representativa", "14/03/2023",
        "Juzgado 1ª Instancia nº 4 Burgos", "412/2023", "No", "Contacto De Prueba",
        "", "Patrimonial y decisiones sanitarias mayores", "Vida diaria", "Sí"]),
]

LEEME = [
    "Plantilla maestra de captura — Personas atendidas · Proyecto ÁGORA",
    "",
    "PARA QUÉ: recoger los datos de las personas de vuestro centro/servicio con la estructura",
    "exacta que ÁGORA importará. Rellenar esta plantilla = carga directa, sin retrabajos.",
    "",
    "CÓMO SE USA:",
    "· Hoja «1 Persona»: UNA fila por persona. Los campos con * son obligatorios.",
    "· Hojas 2-7: varias filas por persona si hace falta (una por alergia, medicamento,",
    "  contacto...). Enlazad cada fila con la persona por su DNI o su código interno.",
    "· La fila gris de cada hoja es un EJEMPLO FICTICIO: borradla antes de entregar.",
    "· Celdas con desplegable: usad las opciones de la lista.",
    "· Fechas: formato dd/mm/aaaa.",
    "",
    "PROTECCIÓN DE DATOS (art. 9 RGPD — datos de salud):",
    "· Este fichero relleno contiene datos reales: SOLO puede vivir en la carpeta",
    "  restringida acordada. Nunca por correo, ni copias locales, ni impresiones.",
    "· Si tenéis duda sobre un campo, dejadlo vacío y anotadlo en notas — mejor vacío que inventado.",
    "",
    "Dudas: Gerencia. Versión de la plantilla: 1.0 · 25/08/2026 (espejo del modelo ÁGORA v0.12+P1).",
]


def generar(salida: Path):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "LEEME"
    ws0.column_dimensions["A"].width = 100
    for i, linea in enumerate(LEEME, 1):
        c = ws0.cell(row=i, column=1, value=linea)
        if i == 1:
            c.font = Font(bold=True, size=14, color="047857")
        c.alignment = Alignment(wrap_text=True)

    for titulo, campos, ejemplo in HOJAS:
        ws = wb.create_sheet(titulo)
        for j, (campo, oblig, _) in enumerate(campos, 1):
            c = ws.cell(row=1, column=j, value=campo)
            c.fill = AMBAR if oblig else VERDE
            c.font = Font(bold=True) if oblig else BLANCO
            ancho = max(16, min(34, len(campo) + 4))
            ws.column_dimensions[get_column_letter(j)].width = ancho
        for j, valor in enumerate(ejemplo, 1):
            c = ws.cell(row=2, column=j, value=valor)
            c.font = GRIS
        for j, (_, _, lista) in enumerate(campos, 1):
            if lista:
                dv = DataValidation(type="list", formula1=lista, allow_blank=True)
                ws.add_data_validation(dv)
                col = get_column_letter(j)
                dv.add(f"{col}3:{col}500")
        ws.freeze_panes = "A2"

    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    print(f"Plantilla generada: {salida}")
    print(f"Hojas: LEEME + {len(HOJAS)} · campos totales: {sum(len(c) for _, c, _ in HOJAS)}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    generar(Path(sys.argv[1]))
