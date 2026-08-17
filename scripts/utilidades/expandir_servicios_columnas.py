"""
expandir_servicios_columnas.py — Convierte "Servicios activos" (multivalor) y
"Centro (Registro PV)" en columnas matriz: una columna por servicio y por
centro, marcadas con X (o ? si el cruce de centro era AMBIGUO).

Uso:
    python scripts/utilidades/expandir_servicios_columnas.py <entrada.xlsx> <salida.xlsx>

Espera el fichero generado por cruzar_ixis_registro_pv.py (21 columnas, con
"Servicios activos" en la col. 19, "Centro (Registro PV)" en la 20 y "Cruce"
en la 21).

RGPD: no imprime datos personales — solo estadísticas agregadas.
"""
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook, Workbook

# Servicios cuyo nombre oficial contiene comas — se protegen antes de separar
SERVICIOS_CON_COMA = [
    "Ocio, cultura y deporte",
    "Promoción, recuperación y mantenimiento de la autonomía",
]

COL_SERVICIOS = 19   # "Servicios activos" (1-index)
COL_CENTRO = 20      # "Centro (Registro PV)"
COL_CRUCE = 21       # "Cruce"


def separar_servicios(valor: str) -> list:
    """Divide el campo multivalor respetando los servicios con coma."""
    texto = valor
    protegidos = {}
    for i, serv in enumerate(SERVICIOS_CON_COMA):
        marca = f"\x00{i}\x00"
        # insensible a mayúsculas para la sustitución
        idx = texto.lower().find(serv.lower())
        while idx != -1:
            texto = texto[:idx] + marca + texto[idx + len(serv):]
            protegidos[marca] = serv
            idx = texto.lower().find(serv.lower())
    partes = []
    for parte in texto.replace(";", ",").split(","):
        p = parte.strip()
        if not p:
            continue
        partes.append(protegidos.get(p, p))
    return partes


def separar_centros(valor: str, cruce: str) -> tuple:
    """Devuelve (lista de centros, marca) según el tipo de cruce."""
    if not valor:
        return [], "X"
    if cruce == "AMBIGUO":
        return [c.strip() for c in valor.split("?") if c.strip()], "?"
    return [c.strip() for c in valor.split("+") if c.strip()], "X"


def expandir(ruta_in: Path, ruta_out: Path) -> None:
    wb = load_workbook(ruta_in, read_only=True, data_only=True)
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    cabecera, datos = list(filas[0]), filas[1:]

    # Pasada 1: inventario de servicios y centros
    cnt_serv, cnt_cen = Counter(), Counter()
    for f in datos:
        serv = str(f[COL_SERVICIOS - 1] or "").strip()
        if serv:
            for s in separar_servicios(serv):
                cnt_serv[s] += 1
        cen = str(f[COL_CENTRO - 1] or "").strip()
        cruce = str(f[COL_CRUCE - 1] or "").strip()
        centros, _ = separar_centros(cen, cruce)
        for c in centros:
            cnt_cen[c] += 1

    servicios = [s for s, _ in cnt_serv.most_common()]
    centros_lista = [c for c, _ in cnt_cen.most_common()]

    # Pasada 2: escribir matriz
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "Ixis matriz servicios-centros"
    ws_out.append(
        cabecera
        + [f"Servicio · {s}" for s in servicios]
        + [f"Centro · {c}" for c in centros_lista]
    )
    for f in datos:
        celdas = list(f)
        serv = str(f[COL_SERVICIOS - 1] or "").strip()
        propios = set(separar_servicios(serv)) if serv else set()
        cen = str(f[COL_CENTRO - 1] or "").strip()
        cruce = str(f[COL_CRUCE - 1] or "").strip()
        cen_propios, marca = separar_centros(cen, cruce)
        fila_serv = ["X" if s in propios else "" for s in servicios]
        fila_cen = [marca if c in set(cen_propios) else "" for c in centros_lista]
        ws_out.append(celdas + fila_serv + fila_cen)

    ruta_out.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(ruta_out)

    print(f"Personas: {len(datos)}")
    print(f"Columnas de servicio: {len(servicios)}")
    for s, n in cnt_serv.most_common():
        print(f"  {n:5d} · Servicio · {s[:60]}")
    print(f"Columnas de centro: {len(centros_lista)}")
    for c, n in cnt_cen.most_common():
        print(f"  {n:5d} · Centro · {c[:60]}")
    print(f"Salida: {ruta_out}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    expandir(Path(sys.argv[1]), Path(sys.argv[2]))
