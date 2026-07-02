"""Prueba en seco del importador 'importar_excel_registro_pv'.

Reproduce la lógica de parseo del comando Django pero SIN tocar la base de
datos: solo lee el Excel y reporta qué se importaría. Útil para validar el
parser antes de que Sistemas Aspanias arranque Docker.

Uso:
    python scripts/prueba_seco_importador.py "ruta/al/Registro PV ... _PSEUDO.xlsx"
"""
from __future__ import annotations

import re
import sys
from datetime import date, datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Falta openpyxl. pip install openpyxl", file=sys.stderr)
    sys.exit(1)


MAPEO_CENTROS = {
    "Res. Quint.": ("RES_QUINT", "Residencia y UA Quintanadueñas", "FAB"),
    "CD Quint.": ("CD_QUINT", "Centro Ocupacional Quintanadueñas", "FAB"),
    "Vicente Aleixandre": ("VICENTE", "Centro Multiactividad Vicente Aleixandre", "FAB"),
    "Puentesauco": ("RES_PUENTESAUCO", "Residencia Puentesauco (RHP)", "Aspaniasmerc"),
    "Salas": ("SAL", "Residencia y Centro de Día Salas", "Aspaniasmerc"),
    "Fuentecillas": ("FUE", "Centro Mayores Fuentecillas", "FAB"),
    "Puentes": ("PUENTES", "Servicio de Vida Independiente", "FAB"),
    "CD Puentesauco": ("CD_PUENTESAUCO", "Centro Ocupacional Puentesauco", "Aspaniasmerc"),
    "Río Arlanza": ("LAR", "Residencia Río Arlanza", "FAB"),
    "Santa Mª": ("SANTA_M", "Residencia Santa María", "FAB"),
    "Viviendas Puentesauco": ("VIV_PUENTESAUCO", "Viviendas Puentesauco", "Aspaniasmerc"),
    "Viviendas Asociación": ("VIV_ASOC", "Viviendas Asociación", "FAB"),
    "Viviendas Fuentecillas": ("VIV_FUE", "Viviendas Fuentecillas", "FAB"),
}

CABECERAS = {
    "usuario/a": "usuario", "usuario": "usuario",
    "gestor/a de caso": "gestor", "gestor de caso": "gestor",
    "persona de referencia": "referencia",
    "pv realizado": "pv_realizado",
    "pv revisado": "pv_revisado",
    "¿está en teams?": "en_teams", "está en teams?": "en_teams", "esta en teams?": "en_teams",
    "¿está en repriss?": "en_repriss", "está en repriss?": "en_repriss", "esta en repriss?": "en_repriss",
    "fecha prevista realización pv": "fecha_prevista",
    "fecha prevista realizacion pv": "fecha_prevista",
    "fecha revisión pv": "fecha_revision", "fecha revision pv": "fecha_revision",
    "estado revisión": "estado_revision", "estado revision": "estado_revision",
    "usuario/a compartido con residencia /vivienda": "compartido",
    "usuario/a compartido con residencia /viv": "compartido",
}


def _norm(s):
    if not isinstance(s, str):
        return ""
    return s.strip().lower()


def _bool_si_no(v):
    if v is None:
        return False
    return str(v).strip().upper() in ("SI", "SÍ", "TRUE", "1", "X")


def _parse_fecha(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    ruta = Path(sys.argv[1])
    if not ruta.exists():
        print(f"ERROR: no existe {ruta}")
        sys.exit(1)

    wb = load_workbook(ruta, data_only=True)
    print(f"Excel: {ruta.name}")
    print(f"Hojas detectadas: {len(wb.sheetnames)}")
    print(f"Mapeo configurado: {len(MAPEO_CENTROS)} centros a importar")
    print()

    resumen_global = {
        "personas_total": 0, "profesionales_unicos": set(),
        "pv_realizados": 0, "pv_revisados": 0, "fuera_plazo": 0,
        "filas_invalidas": 0,
    }

    for hoja_nombre, (codigo, nombre_canonico, corresponsable) in MAPEO_CENTROS.items():
        if hoja_nombre not in wb.sheetnames:
            print(f"  [SIN HOJA] '{hoja_nombre}' no encontrada en el Excel")
            continue
        ws = wb[hoja_nombre]

        # Localizar fila de cabecera
        fila_cab = None
        for r_idx in range(1, 15):
            row_vals = [_norm(c.value) for c in ws[r_idx]]
            if "usuario/a" in row_vals or "usuario" in row_vals:
                fila_cab = r_idx
                break

        if not fila_cab:
            print(f"  [SIN CABECERA] '{hoja_nombre}': no se localiza fila 'Usuario/a'")
            continue

        # Indexar columnas
        col_idx = {}
        for c_idx, cell in enumerate(ws[fila_cab], start=1):
            clave = CABECERAS.get(_norm(cell.value))
            if clave:
                col_idx[clave] = c_idx

        # Procesar filas
        n_personas = 0
        n_realizados = 0
        n_revisados = 0
        n_fuera_plazo = 0
        n_invalidas = 0
        profesionales_hoja = set()

        for row in ws.iter_rows(min_row=fila_cab + 1, values_only=True):
            def val(clave):
                idx = col_idx.get(clave)
                if idx is None or idx - 1 >= len(row):
                    return None
                return row[idx - 1]

            nombre_usuario = val("usuario")
            if not nombre_usuario or not isinstance(nombre_usuario, str) or not nombre_usuario.strip():
                continue
            partes = nombre_usuario.strip().split()
            if len(partes) < 2:
                n_invalidas += 1
                continue

            n_personas += 1
            if _bool_si_no(val("pv_realizado")):
                n_realizados += 1
            if _bool_si_no(val("pv_revisado")):
                n_revisados += 1
            estado_rev = (val("estado_revision") or "").strip().lower() if val("estado_revision") else ""
            if "fuera" in estado_rev:
                n_fuera_plazo += 1
            for clave_prof in ("gestor", "referencia"):
                v = val(clave_prof)
                if v and isinstance(v, str) and v.strip():
                    profesionales_hoja.add(v.strip())

        resumen_global["personas_total"] += n_personas
        resumen_global["pv_realizados"] += n_realizados
        resumen_global["pv_revisados"] += n_revisados
        resumen_global["fuera_plazo"] += n_fuera_plazo
        resumen_global["filas_invalidas"] += n_invalidas
        resumen_global["profesionales_unicos"].update(profesionales_hoja)

        print(f"  [{codigo:>16}] {hoja_nombre:<25} ({corresponsable:>12}) | "
              f"{n_personas:>3} personas | "
              f"{n_realizados:>3} PV realizados | "
              f"{n_revisados:>3} revisados | "
              f"{n_fuera_plazo:>2} fuera plazo | "
              f"{len(profesionales_hoja):>2} profs · "
              f"{n_invalidas} filas invalidas")

    print()
    print("=" * 80)
    print("RESUMEN GLOBAL")
    print("=" * 80)
    print(f"  Personas atendidas que se importarían:   {resumen_global['personas_total']}")
    print(f"  Profesionales únicos que se crearían:    {len(resumen_global['profesionales_unicos'])}")
    print(f"  PV realizados:                            {resumen_global['pv_realizados']}")
    print(f"  PV revisados:                             {resumen_global['pv_revisados']}")
    print(f"  PV fuera de plazo:                        {resumen_global['fuera_plazo']}")
    print(f"  Filas inválidas (nombre incompleto):      {resumen_global['filas_invalidas']}")
    print()
    print("Comparativa con la hoja 'Datos Grupo Aspanias' del Excel:")
    print("  Total usuarios Excel:    383")
    print(f"  Total usuarios parser:   {resumen_global['personas_total']}")
    print()
    if resumen_global['personas_total'] >= 380 and resumen_global['personas_total'] <= 400:
        print("  -> Parser coherente con los datos globales del Excel.")
    else:
        print("  -> Revisar diferencia.")


if __name__ == "__main__":
    main()
