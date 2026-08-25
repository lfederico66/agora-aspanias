"""
preparar_revisiones.py — Genera los ficheros de revisión para los equipos a
partir de los productos de migración:

  1) Para Dirección Fuentecillas (Laura): censo con columnas de respuesta
     y filas problemáticas resaltadas (cruce ambiguo · modalidad dudosa).
  2) Para Trabajo Social (Uxue): maestro con columnas de resolución en la
     hoja Excepciones (desplegable de motivo) y en los conflictos.

Uso:
    python scripts/utilidades/preparar_revisiones.py \
        <Fuentecillas_maestro.xlsx> <MAESTRO_importacion_agora.xlsx> <carpeta_salida>

RGPD: no imprime datos personales — solo estadísticas. La salida debe ir a
una ubicación de custodia acordada (carpeta restringida).
"""
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AMBAR = PatternFill("solid", fgColor="FEF3C7")
AZUL = PatternFill("solid", fgColor="DBEAFE")
VERDE_CAB = PatternFill("solid", fgColor="047857")
BLANCO_NEGRITA = Font(color="FFFFFF", bold=True)
RESPUESTA = PatternFill("solid", fgColor="ECFDF5")


def hoja_instrucciones(wb, titulo, lineas):
    ws = wb.create_sheet("LEEME", 0)
    ws.column_dimensions["A"].width = 110
    c = ws.cell(row=1, column=1, value=titulo)
    c.font = Font(bold=True, size=14, color="047857")
    for i, linea in enumerate(lineas, 3):
        celda = ws.cell(row=i, column=1, value=linea)
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def estilo_cabecera(ws, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.fill = VERDE_CAB
        c.font = BLANCO_NEGRITA
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{ws.max_row}"


def preparar_laura(ruta_censo: Path, salida: Path):
    src = load_workbook(ruta_censo, read_only=True)
    filas = list(src.active.iter_rows(values_only=True))
    cab, datos = list(filas[0]), filas[1:]

    wb = Workbook()
    ws = wb.active
    ws.title = "Censo Fuentecillas"
    extra = ["RESPUESTA: ¿Centro correcto?", "RESPUESTA: Observaciones"]
    ws.append(cab + extra)
    idx_cruce = cab.index("Cruce")
    idx_mod = cab.index("Modalidad")

    ambiguos = dudosos = 0
    for f in datos:
        ws.append(list(f) + ["", ""])
        fila_x = ws.max_row
        if str(f[idx_cruce] or "").upper().startswith("AMBIGUO"):
            ambiguos += 1
            for col in range(1, len(cab) + 1):
                ws.cell(row=fila_x, column=col).fill = AMBAR
        elif f[idx_mod] == "Otro servicio":
            dudosos += 1
            for col in range(1, len(cab) + 1):
                ws.cell(row=fila_x, column=col).fill = AZUL
    for col_extra in (len(cab) + 1, len(cab) + 2):
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col_extra).fill = RESPUESTA
        ws.column_dimensions[get_column_letter(col_extra)].width = 30
    estilo_cabecera(ws, len(cab) + len(extra))

    hoja_instrucciones(wb, "Revisión del censo de Fuentecillas — instrucciones", [
        "Hola Laura. Este es el censo que saldrá de la migración de datos a ÁGORA (pestaña siguiente).",
        "",
        "QUÉ REVISAR (dos filtros, columnas de respuesta en verde al final):",
        "",
        "1 · FILAS EN ÁMBAR (cruce AMBIGUO): el nombre corto del Registro PV coincide con más de una persona. "
        "En la columna «Centro (Registro PV)» verás los centros candidatos separados por «?». "
        "Escribe en «RESPUESTA: ¿Centro correcto?» el centro que corresponde (o «No es de Fuentecillas»).",
        "",
        "2 · FILAS EN AZUL (modalidad «Otro servicio»): figuran en la hoja Fuentecillas del Registro PV "
        "pero su ficha de Ixis no marca ni Residencia ni Centro de Día. Confirma en «RESPUESTA» si son del "
        "centro (y en qué modalidad) o si es un desfase de altas/bajas.",
        "",
        "El resto de filas no necesitan nada (puedes anotar observaciones si ves algo raro).",
        "Cuando termines, avisa a Gerencia. Objetivo: antes del 5 de septiembre.",
        "",
        "Este fichero contiene datos reales — no lo reenvíes ni lo saques de la carpeta compartida.",
    ])
    wb.save(salida)
    return len(datos), ambiguos, dudosos


def preparar_uxue(ruta_maestro: Path, salida: Path):
    src = load_workbook(ruta_maestro, read_only=True)

    wb = Workbook()
    wb.remove(wb.active)

    # --- Hoja Excepciones con columnas de resolución ---
    ws_exc_src = src["Excepciones"]
    filas = list(ws_exc_src.iter_rows(values_only=True))
    ws = wb.create_sheet("Excepciones")
    extra = ["RESPUESTA: Motivo real", "RESPUESTA: Corrección / observación"]
    ws.append(list(filas[0]) + extra)
    for f in filas[1:]:
        ws.append(list(f) + ["", ""])
    ncols = len(filas[0])
    dv = DataValidation(
        type="list",
        formula1='"Errata de nombre,Baja de la entidad,Alta nueva (no está en Ixis),Duplicado,Otro"',
        allow_blank=True,
    )
    ws.add_data_validation(dv)
    col_dv = get_column_letter(ncols + 1)
    dv.add(f"{col_dv}2:{col_dv}{ws.max_row}")
    for col_extra in (ncols + 1, ncols + 2):
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col_extra).fill = RESPUESTA
        ws.column_dimensions[get_column_letter(col_extra)].width = 32
    estilo_cabecera(ws, ncols + len(extra))
    n_exc = ws.max_row - 1

    # --- Hoja Conflictos: solo las filas con conflicto, con columna de decisión ---
    ws_m = src["Maestro"]
    filas_m = list(ws_m.iter_rows(values_only=True))
    cab_m = list(filas_m[0])
    idx_cfn = cab_m.index("Conflicto F.Nacimiento")
    idx_cdis = cab_m.index("Conflicto % Disc")
    ws2 = wb.create_sheet("Conflictos")
    ws2.append(cab_m + ["RESPUESTA: valor correcto", "RESPUESTA: observación"])
    n_conf = 0
    for f in filas_m[1:]:
        if f[idx_cfn] == "SÍ" or f[idx_cdis] == "SÍ":
            n_conf += 1
            ws2.append(list(f) + ["", ""])
            for col in range(1, len(cab_m) + 1):
                ws2.cell(row=ws2.max_row, column=col).fill = AMBAR
    for col_extra in (len(cab_m) + 1, len(cab_m) + 2):
        for r in range(2, ws2.max_row + 1):
            ws2.cell(row=r, column=col_extra).fill = RESPUESTA
        ws2.column_dimensions[get_column_letter(col_extra)].width = 30
    estilo_cabecera(ws2, len(cab_m) + 2)

    hoja_instrucciones(wb, "Revisión de excepciones y conflictos de la migración — instrucciones", [
        "Hola Uxue. Del cruce de vuestros ficheros (DNI/TSS, diagnósticos, contactos) con el volcado de "
        "Ixis quedan casos que necesitan criterio de Trabajo Social. Dos pestañas:",
        "",
        "1 · EXCEPCIONES: personas de vuestros ficheros que no casan con ninguna del Ixis. "
        "Elige el «Motivo real» en el desplegable (errata de nombre, baja, alta nueva, duplicado, otro) "
        "y, si es errata, escribe el nombre correcto en la columna de corrección.",
        "",
        "2 · CONFLICTOS: personas donde dos fuentes dicen cosas distintas (fecha de nacimiento o "
        "% de discapacidad — mira las columnas «Conflicto...» en SÍ). Escribe en «RESPUESTA: valor "
        "correcto» el dato bueno (el del documento oficial más reciente).",
        "",
        "Cuando termines, avisa a Gerencia. Objetivo: antes del 5 de septiembre.",
        "",
        "Este fichero contiene datos reales — no lo reenvíes ni lo saques de la carpeta compartida.",
    ])
    wb.save(salida)
    return n_exc, n_conf


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    censo, maestro, carpeta = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
    carpeta.mkdir(parents=True, exist_ok=True)

    f1 = carpeta / "REVISION Fuentecillas (Laura).xlsx"
    total, amb, dud = preparar_laura(censo, f1)
    print(f"Laura: {total} personas · {amb} ambiguas (ámbar) · {dud} dudosas (azul) → {f1.name}")

    f2 = carpeta / "REVISION Migracion (Trabajo Social).xlsx"
    n_exc, n_conf = preparar_uxue(maestro, f2)
    print(f"Uxue: {n_exc} excepciones · {n_conf} conflictos → {f2.name}")
    print(f"Carpeta: {carpeta}")
