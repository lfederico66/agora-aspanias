"""Vistas del Plan de Vida (Fundación Aspanias)."""
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.shortcuts import get_object_or_404, redirect, render

from personas.models import Centro, PersonaAtendida

from .forms import (
    AlgoSobreMiForm,
    EntradaPlanApoyoFormSet,
    HistoriaDeVidaForm,
    PlanDeApoyoForm,
    ProyectoDeVidaForm,
    RespuestaHistoriaVidaFormSet,
    RevisionObjetivoForm,
    SeccionAlgoSobreMiFormSet,
)
from .models import (
    CLAVES_RELACION_HISTORIA_VIDA,
    PREGUNTAS_HISTORIA_VIDA,
    SECCIONES_ALGO_SOBRE_MI,
    AlgoSobreMi,
    DocumentoPlanDeVida,
    HistoriaDeVida,
    PlanDeApoyo,
    PlanDeVida,
    ProyectoDeVida,
    RespuestaHistoriaVida,
    RevisionObjetivo,
    SeccionAlgoSobreMi,
)


@login_required
def detalle_plan_vida(request, persona_id):
    """Muestra el Plan de Vida vigente de la persona, o el último si no hay vigente."""
    persona = get_object_or_404(PersonaAtendida, pk=persona_id)
    plan = (
        persona.planes_vida
        .filter(estado=PlanDeVida.Estado.VIGENTE)
        .first()
        or persona.planes_vida.order_by("-anualidad").first()
    )

    documentos_por_tipo = {}
    objetivos_por_ambito = {}
    cambios_recientes = []
    if plan:
        existentes = {d.tipo: d for d in plan.documentos.all()}
        for tipo, etiqueta in DocumentoPlanDeVida.Tipo.choices:
            documentos_por_tipo[etiqueta] = existentes.get(tipo)
        for obj in plan.objetivos.all():
            objetivos_por_ambito.setdefault(obj.get_ambito_display(), []).append(obj)
        cambios_recientes = persona.cambios_significativos.order_by("-fecha_deteccion")[:10]

    return render(request, "pia/detalle.html", {
        "persona": persona,
        "plan": plan,
        "documentos_por_tipo": documentos_por_tipo,
        "objetivos_por_ambito": objetivos_por_ambito,
        "cambios_recientes": cambios_recientes,
        "planes_historicos": persona.planes_vida.exclude(pk=plan.pk if plan else None).order_by("-anualidad")[:5] if plan else [],
    })


@login_required
def panel_objetivos(request, persona_id):
    """Panel imprimible (A4) con los objetivos visibles del Plan de Vida.

    Protocolo §6: los objetivos pueden reflejarse en paneles visuales / corchos
    del centro. Solo se muestran los objetivos con `visible_en_panel=True` y si
    la persona ha autorizado la visibilidad (`plan.visible_en_paneles=True`).
    """
    persona = get_object_or_404(PersonaAtendida, pk=persona_id)
    plan = persona.planes_vida.filter(estado=PlanDeVida.Estado.VIGENTE).first()
    objetivos = []
    if plan and plan.visible_en_paneles:
        objetivos = list(
            plan.objetivos.filter(visible_en_panel=True).select_related("responsable_seguimiento")
        )
    return render(request, "pia/panel_imprimible.html", {
        "persona": persona,
        "plan": plan,
        "objetivos": objetivos,
        "autorizado": bool(plan and plan.visible_en_paneles),
    })


@login_required
def documento_revision_objetivos(request, persona_id, anualidad):
    """Vista imprimible del documento 'Revisión de objetivos' (doc 3 del protocolo).

    Replica exactamente el layout del Excel actual (4 columnas) — para sustituir
    nativamente la plantilla compartida.
    """
    persona = get_object_or_404(PersonaAtendida, pk=persona_id)
    plan = get_object_or_404(PlanDeVida, persona=persona, anualidad=anualidad)

    objetivos = list(plan.objetivos.select_related("responsable_seguimiento"))
    revisiones_por_objetivo = {
        r.objetivo_id: r
        for r in RevisionObjetivo.objects.filter(
            objetivo__plan_vida=plan, anualidad=anualidad,
        ).select_related("realizada_por")
    }
    filas = [
        {"objetivo": obj, "revision": revisiones_por_objetivo.get(obj.id)}
        for obj in objetivos
    ]
    return render(request, "pia/revision_objetivos.html", {
        "persona": persona,
        "plan": plan,
        "filas": filas,
        "anualidad": anualidad,
    })


def _resolver_plan(persona_id, anualidad):
    persona = get_object_or_404(PersonaAtendida, pk=persona_id)
    plan = get_object_or_404(PlanDeVida, persona=persona, anualidad=anualidad)
    return persona, plan


@login_required
def documento_historia_vida(request, persona_id, anualidad):
    persona, plan = _resolver_plan(persona_id, anualidad)
    historia = getattr(plan, "historia_vida", None)
    respuestas = historia.respuestas.all() if historia else []
    return render(request, "pia/historia_vida.html", {
        "persona": persona, "plan": plan, "historia": historia,
        "respuestas": respuestas,
        "claves_metodologicas": CLAVES_RELACION_HISTORIA_VIDA,
    })


@login_required
def documento_algo_sobre_mi(request, persona_id, anualidad):
    persona, plan = _resolver_plan(persona_id, anualidad)
    algo = getattr(plan, "algo_sobre_mi", None)
    secciones = algo.secciones.all() if algo else []
    return render(request, "pia/algo_sobre_mi.html", {
        "persona": persona, "plan": plan, "algo": algo, "secciones": secciones,
    })


@login_required
def documento_plan_apoyo(request, persona_id, anualidad):
    persona, plan = _resolver_plan(persona_id, anualidad)
    plan_apoyo = getattr(plan, "plan_apoyo", None)
    entradas = []
    if plan_apoyo:
        entradas = plan_apoyo.entradas.prefetch_related("dimensiones_cdv").all()
    return render(request, "pia/plan_apoyo.html", {
        "persona": persona, "plan": plan, "plan_apoyo": plan_apoyo, "entradas": entradas,
    })


@login_required
def documento_proyecto_vida(request, persona_id, anualidad):
    persona, plan = _resolver_plan(persona_id, anualidad)
    proyecto = getattr(plan, "proyecto_vida", None)
    return render(request, "pia/proyecto_vida.html", {
        "persona": persona, "plan": plan, "proyecto": proyecto,
    })


# ---------------------------------------------------------------------------
# EDICIÓN web (M2)
# ---------------------------------------------------------------------------


def _get_or_create_historia(plan):
    historia, creado = HistoriaDeVida.objects.get_or_create(plan_vida=plan)
    if creado or not historia.respuestas.exists():
        for idx, (codigo, pregunta) in enumerate(PREGUNTAS_HISTORIA_VIDA, start=1):
            RespuestaHistoriaVida.objects.get_or_create(
                historia=historia, pregunta_codigo=codigo,
                defaults={"pregunta_texto": pregunta, "orden": idx},
            )
    return historia


def _get_or_create_algo(plan):
    algo, creado = AlgoSobreMi.objects.get_or_create(plan_vida=plan)
    if creado or not algo.secciones.exists():
        for idx, (codigo, titulo) in enumerate(SECCIONES_ALGO_SOBRE_MI, start=1):
            SeccionAlgoSobreMi.objects.get_or_create(
                documento=algo, codigo=codigo,
                defaults={"titulo": titulo, "orden": idx},
            )
    return algo


@login_required
@transaction.atomic
def editar_historia_vida(request, persona_id, anualidad):
    persona, plan = _resolver_plan(persona_id, anualidad)
    historia = _get_or_create_historia(plan)
    if request.method == "POST":
        form = HistoriaDeVidaForm(request.POST, instance=historia)
        formset = RespuestaHistoriaVidaFormSet(request.POST, instance=historia)
        if form.is_valid() and formset.is_valid():
            form.save()
            formset.save()
            messages.success(request, "Historia de Vida guardada.")
            return redirect("planes_vida:historia_vida", persona_id=persona.id, anualidad=anualidad)
    else:
        form = HistoriaDeVidaForm(instance=historia)
        formset = RespuestaHistoriaVidaFormSet(instance=historia)
    return render(request, "pia/editar_historia_vida.html", {
        "persona": persona, "plan": plan, "historia": historia,
        "form": form, "formset": formset,
        "claves_metodologicas": CLAVES_RELACION_HISTORIA_VIDA,
    })


@login_required
@transaction.atomic
def editar_algo_sobre_mi(request, persona_id, anualidad):
    persona, plan = _resolver_plan(persona_id, anualidad)
    algo = _get_or_create_algo(plan)
    if request.method == "POST":
        form = AlgoSobreMiForm(request.POST, instance=algo)
        formset = SeccionAlgoSobreMiFormSet(request.POST, instance=algo)
        if form.is_valid() and formset.is_valid():
            form.save()
            formset.save()
            messages.success(request, "Algo sobre mí guardado.")
            return redirect("planes_vida:algo_sobre_mi", persona_id=persona.id, anualidad=anualidad)
    else:
        form = AlgoSobreMiForm(instance=algo)
        formset = SeccionAlgoSobreMiFormSet(instance=algo)
    return render(request, "pia/editar_algo_sobre_mi.html", {
        "persona": persona, "plan": plan, "algo": algo,
        "form": form, "formset": formset,
    })


@login_required
@transaction.atomic
def editar_revision_objetivos(request, persona_id, anualidad):
    """Permite editar la RevisionObjetivo de cada objetivo del plan."""
    persona, plan = _resolver_plan(persona_id, anualidad)
    objetivos = list(plan.objetivos.all())

    if request.method == "POST":
        ok = True
        for obj in objetivos:
            revision, _ = RevisionObjetivo.objects.get_or_create(
                objetivo=obj, anualidad=anualidad,
                defaults={"fecha_revision": plan.fecha_revision or plan.fecha_elaboracion},
            )
            form = RevisionObjetivoForm(request.POST, prefix=f"obj_{obj.id}", instance=revision)
            if form.is_valid():
                form.save()
            else:
                ok = False
        if ok:
            messages.success(request, "Revisiones guardadas.")
            return redirect("planes_vida:revision_objetivos", persona_id=persona.id, anualidad=anualidad)

    formularios = []
    for obj in objetivos:
        revision = RevisionObjetivo.objects.filter(objetivo=obj, anualidad=anualidad).first()
        if not revision:
            revision = RevisionObjetivo(
                objetivo=obj, anualidad=anualidad,
                fecha_revision=plan.fecha_revision or plan.fecha_elaboracion,
            )
        formularios.append({
            "objetivo": obj,
            "form": RevisionObjetivoForm(prefix=f"obj_{obj.id}", instance=revision),
        })
    return render(request, "pia/editar_revision_objetivos.html", {
        "persona": persona, "plan": plan, "anualidad": anualidad,
        "formularios": formularios,
    })


@login_required
@transaction.atomic
def editar_plan_apoyo(request, persona_id, anualidad):
    persona, plan = _resolver_plan(persona_id, anualidad)
    plan_apoyo, _ = PlanDeApoyo.objects.get_or_create(plan_vida=plan)
    if request.method == "POST":
        form = PlanDeApoyoForm(request.POST, instance=plan_apoyo)
        formset = EntradaPlanApoyoFormSet(request.POST, instance=plan_apoyo)
        if form.is_valid() and formset.is_valid():
            form.save()
            formset.save()
            messages.success(request, "Plan de Apoyo guardado.")
            return redirect("planes_vida:plan_apoyo", persona_id=persona.id, anualidad=anualidad)
    else:
        form = PlanDeApoyoForm(instance=plan_apoyo)
        formset = EntradaPlanApoyoFormSet(instance=plan_apoyo)
    return render(request, "pia/editar_plan_apoyo.html", {
        "persona": persona, "plan": plan, "plan_apoyo": plan_apoyo,
        "form": form, "formset": formset,
    })


@login_required
@transaction.atomic
def editar_proyecto_vida(request, persona_id, anualidad):
    persona, plan = _resolver_plan(persona_id, anualidad)
    proyecto, _ = ProyectoDeVida.objects.get_or_create(plan_vida=plan)
    if request.method == "POST":
        form = ProyectoDeVidaForm(request.POST, instance=proyecto)
        if form.is_valid():
            form.save()
            messages.success(request, "Proyecto de Vida guardado.")
            return redirect("planes_vida:proyecto_vida", persona_id=persona.id, anualidad=anualidad)
    else:
        form = ProyectoDeVidaForm(instance=proyecto)
    return render(request, "pia/editar_proyecto_vida.html", {
        "persona": persona, "plan": plan, "proyecto": proyecto,
        "form": form,
    })


# ---------------------------------------------------------------------------
# Control de Planes de Vida por servicios (sustituye al Excel global)
# ---------------------------------------------------------------------------


@login_required
def control_pv_por_servicio(request, centro_codigo=None):
    """Reproduce el 'Registro de Control de Proyectos de Vida por Servicio'.

    Si no se indica centro, lista todos los centros con un resumen para elegir.
    """
    centros = Centro.objects.filter(activo=True).order_by("nombre")

    if centro_codigo is None:
        # Tablero de selección — resumen por centro
        resumen = []
        agregado = {
            "personas": 0, "realizados": 0, "pendientes": 0,
            "revisados": 0, "fuera_plazo": 0,
        }
        for c in centros:
            personas_activas = c.personas.filter(fecha_baja__isnull=True).count()
            planes_qs = PlanDeVida.objects.filter(
                persona__centro_referencia=c,
                persona__fecha_baja__isnull=True,
            )
            planes_realizados = planes_qs.filter(pv_realizado=True).count()
            planes_pendientes = planes_qs.filter(pv_realizado=False).count()
            planes_revisados = planes_qs.filter(pv_revisado=True).count()
            fuera_plazo = planes_qs.filter(estado_revision="fuera_plazo").count()
            pct_realizado = round(100 * planes_realizados / personas_activas, 1) if personas_activas else 0

            # Estado visual: al día (≥90%), atrasado (75-89%), crítico (<75%)
            if pct_realizado >= 90:
                estado_visual = "al_dia"
            elif pct_realizado >= 75:
                estado_visual = "atrasado"
            else:
                estado_visual = "critico"

            resumen.append({
                "centro": c,
                "personas": personas_activas,
                "realizados": planes_realizados,
                "pendientes": planes_pendientes,
                "revisados": planes_revisados,
                "fuera_plazo": fuera_plazo,
                "pct_realizado": pct_realizado,
                "estado_visual": estado_visual,
            })
            agregado["personas"] += personas_activas
            agregado["realizados"] += planes_realizados
            agregado["pendientes"] += planes_pendientes
            agregado["revisados"] += planes_revisados
            agregado["fuera_plazo"] += fuera_plazo

        agregado["pct_realizado"] = (
            round(100 * agregado["realizados"] / agregado["personas"], 1)
            if agregado["personas"] else 0
        )
        return render(request, "pia/control_pv_centros.html", {
            "resumen": resumen,
            "agregado": agregado,
        })

    centro = get_object_or_404(Centro, codigo=centro_codigo)
    personas = (
        centro.personas
        .filter(fecha_baja__isnull=True)
        .select_related("gestor_caso", "persona_referencia")
        .order_by("apellido_1", "apellido_2", "nombre")
    )
    # Para cada persona, el plan vigente (o el más reciente)
    filas = []
    for idx, p in enumerate(personas, start=1):
        plan = p.planes_vida.filter(estado=PlanDeVida.Estado.VIGENTE).first() or \
               p.planes_vida.order_by("-anualidad").first()
        documentos_en_teams = False
        documentos_en_repriss = False
        if plan:
            docs = list(plan.documentos.all())
            documentos_en_teams = any(d.url_almacenamiento for d in docs)
            documentos_en_repriss = any(d.publicado_en_repriss for d in docs)
        filas.append({
            "n": idx,
            "persona": p,
            "gestor": p.gestor_caso,
            "referencia": p.persona_referencia,
            "plan": plan,
            "en_teams": documentos_en_teams,
            "en_repriss": documentos_en_repriss,
        })
    return render(request, "pia/control_pv_servicio.html", {
        "centro": centro, "filas": filas, "n_personas": len(filas),
    })


@login_required
def exportar_control_pv_excel(request, centro_codigo):
    """Genera el .xlsx con el formato exacto del 'Registro PV por servicios'.

    Útil para periodo de transición — los responsables que aún quieran su Excel
    de siempre pueden descargarlo desde ÁGORA.
    """
    from io import BytesIO

    from django.http import HttpResponse
    from django.utils import timezone
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    centro = get_object_or_404(Centro, codigo=centro_codigo)
    personas = (
        centro.personas
        .filter(fecha_baja__isnull=True)
        .select_related("gestor_caso", "persona_referencia")
        .order_by("apellido_1", "apellido_2", "nombre")
    )

    wb = Workbook()
    ws = wb.active
    ws.title = centro.codigo[:31]

    # Cabecera del documento (filas 2-5) — replica el Excel original
    ws.cell(row=2, column=2, value="REGISTRO DE CONTROL DE PROYECTOS DE VIDA POR SERVICIO")
    ws.cell(row=2, column=2).font = Font(bold=True, size=14)
    ws.cell(row=4, column=2, value="Centro/Servicio:")
    ws.cell(row=4, column=3, value=centro.nombre)
    ws.cell(row=5, column=2, value="Responsable:")
    # El responsable principal: gestor del primer plan vigente del centro, si lo hay
    primer_plan = PlanDeVida.objects.filter(
        persona__centro_referencia=centro,
        estado=PlanDeVida.Estado.VIGENTE,
    ).select_related("gestor_caso").first()
    responsable = primer_plan.gestor_caso.nombre_completo if primer_plan and primer_plan.gestor_caso else ""
    ws.cell(row=5, column=3, value=responsable)
    ws.cell(row=5, column=10, value="Fecha")
    ws.cell(row=5, column=11, value=timezone.localdate().strftime("%d/%m/%Y"))

    # Cabecera de tabla (fila 7)
    cabeceras = [
        "#", "Usuario/a", "Gestor/a de caso", "Persona de referencia",
        "PV realizado", "PV revisado", "¿Está en Teams?", "¿Está en Repriss?",
        "Fecha prevista realización PV", "Fecha revisión PV", "Estado revisión",
        "Usuario/a compartido con residencia /vivienda",
    ]
    fill = PatternFill(start_color="1E6B4B", end_color="1E6B4B", fill_type="solid")
    for c_idx, texto in enumerate(cabeceras, start=1):
        celda = ws.cell(row=7, column=c_idx, value=texto)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = fill
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Datos (fila 8 en adelante)
    for idx, p in enumerate(personas, start=1):
        plan = p.planes_vida.filter(estado=PlanDeVida.Estado.VIGENTE).first() or \
               p.planes_vida.order_by("-anualidad").first()
        documentos_en_teams = False
        documentos_en_repriss = False
        if plan:
            docs = list(plan.documentos.all())
            documentos_en_teams = any(d.url_almacenamiento for d in docs)
            documentos_en_repriss = any(d.publicado_en_repriss for d in docs)

        fila = 7 + idx
        valores = [
            idx,
            f"{p.apellido_1} {p.apellido_2}, {p.nombre}".strip(),
            p.gestor_caso.nombre_completo if p.gestor_caso else "",
            p.persona_referencia.nombre_completo if p.persona_referencia else "",
            "SI" if (plan and plan.pv_realizado) else "NO",
            "SI" if (plan and plan.pv_revisado) else "NO",
            "SI" if documentos_en_teams else "NO",
            "SI" if documentos_en_repriss else "NO",
            plan.fecha_prevista_realizacion.strftime("%d/%m/%Y") if plan and plan.fecha_prevista_realizacion else "",
            plan.fecha_revision.strftime("%d/%m/%Y") if plan and plan.fecha_revision else "",
            plan.get_estado_revision_display() if plan else "",
            "SI" if (plan and plan.compartido_con_residencia_vivienda) else "NO",
        ]
        for c_idx, v in enumerate(valores, start=1):
            ws.cell(row=fila, column=c_idx, value=v)

    # Anchos de columna
    anchos = [4, 32, 26, 26, 12, 12, 14, 14, 22, 18, 16, 22]
    for c_idx, w in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=7, column=c_idx).column_letter].width = w

    # Producir respuesta
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    nombre_archivo = f"Control_PV_{centro.codigo}_{timezone.localdate():%Y%m%d}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{nombre_archivo}"'
    return response
