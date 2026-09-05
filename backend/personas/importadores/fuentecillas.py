"""
Importador Fuentecillas — capa pura (sin ORM).

Lee el fichero consolidado de personas atendidas de Fuentecillas, normaliza,
valida y devuelve registros listos para persistir + excepciones. Ejecutable
en seco sin Django ni base de datos:

    python backend/personas/importadores/fuentecillas.py <fichero.xlsx> <informe_salida.xlsx>

Decisiones de mapeo (validar con Dirección/DPO):
  - Teléfonos 1-2 y mail se asignan al CONTACTO principal (así venían en Ixis);
    la dirección postal, a la persona.
  - Medicación habitual, intervenciones quirúrgicas, ICAP y expediente entran
    como texto migrado en `notas_relevantes` — el detalle estructurado (pautas)
    lo construirá Enfermería durante el pre-piloto.
  - «Ingresos mensuales» NO se importa (minimización, art. 5 RGPD: sin
    finalidad documentada en el pre-piloto).
La consola solo muestra estadísticas; los nombres van únicamente al informe.
"""
from __future__ import annotations

import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

PARTICULAS = {"DE", "DEL", "LA", "LAS", "LOS", "SAN", "SANTA", "Y"}
LETRAS_NIF = "TRWAGMYFPDXBNJZSQVHLCKE"

CENTROS = {
    "RESIDENCIA": {"codigo": "FUE-MAY", "nombre": "Centro de Mayores Fuentecillas — Residencia"},
    "CD": {"codigo": "FUE-CDM", "nombre": "Centro de Mayores Fuentecillas — Centro de Día"},
}

RELACIONES = {
    "MADRE": "madre", "PADRE": "padre", "HERMANO": "hermano", "HERMANA": "hermano",
    "HIJO": "hijo", "HIJA": "hijo", "SOBRINO": "otro", "SOBRINA": "otro",
    "CONYUGE": "conyuge", "ESPOSO": "conyuge", "ESPOSA": "conyuge", "PAREJA": "conyuge",
    "TUTOR": "tutor", "TUTORA": "tutor", "CURADOR": "curador", "CURADORA": "curador",
    "AMIGO": "amistad", "AMIGA": "amistad", "AMISTAD": "amistad",
    "FUNDACION": "curador", "ENTIDAD": "curador",
}

COLS = {  # índice 0-based → campo
    "apellidos": 0, "nombre": 1, "edad": 2, "parentesco": 3, "contacto_nombre": 4,
    "direccion": 5, "cp": 6, "ciudad": 7, "provincia": 8, "tel1": 9, "tel2": 10,
    "mail": 11, "fecha_cambio_med": 12, "medicacion": 13, "enf_cronicas": 14,
    "quirurgicas": 15, "prestaciones": 16, "dni": 17, "cad_dni": 18, "ingresos": 19,
    "centro": 20, "f_nac": 21, "f_res_dep": 22, "grado_dep": 23, "pct_disc": 24,
    "bvd": 25, "icap": 26, "f_icap": 27, "tsi": 28, "nuss": 29, "expediente": 30,
    "diag_1": 31, "diag_2": 32, "f_validez_diag": 33, "pd": 34, "pfs": 35,
    "pmr": 36, "p3p": 37, "forma_com": 38, "f_incorporacion": 39,
}


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return re.sub(r"\s+", " ", str(v)).strip()


def _sin_tildes(t: str) -> str:
    d = unicodedata.normalize("NFD", t)
    return "".join(c for c in d if unicodedata.category(c) != "Mn")


def normalizar_fecha(v):
    """→ date | None."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = _texto(v).split()[0] if _texto(v) else ""
    for sep in ("/", "-", "."):
        p = t.split(sep)
        if len(p) == 3 and all(x.isdigit() for x in p):
            try:
                if len(p[0]) == 4:
                    return date(int(p[0]), int(p[1]), int(p[2]))
                return date(int(p[2]), int(p[1]), int(p[0]))
            except ValueError:
                return None
    return None


def validar_nif(v) -> bool:
    t = "".join(c for c in _texto(v).upper() if c.isalnum())
    if len(t) != 9:
        return False
    num = t[:8]
    if t[0] in "XYZ":
        num = {"X": "0", "Y": "1", "Z": "2"}[t[0]] + t[1:8]
    return num.isdigit() and t[8] == LETRAS_NIF[int(num) % 23]


def separar_apellidos(apellidos: str) -> tuple[str, str]:
    partes = apellidos.split()
    if not partes:
        return "", ""
    if len(partes) == 1:
        return partes[0], ""
    # agrupar partículas con la palabra siguiente (DE LA FUENTE PEREZ → [DE LA FUENTE, PEREZ])
    grupos, actual = [], []
    for p in partes:
        actual.append(p)
        if _sin_tildes(p.upper()) not in PARTICULAS:
            grupos.append(" ".join(actual))
            actual = []
    if actual:
        grupos.append(" ".join(actual))
    if len(grupos) == 1:
        return grupos[0], ""
    return grupos[0], " ".join(grupos[1:])


def normalizar_centro(v: str):
    """→ ('RESIDENCIA'|'CD', None) o (None, motivo_excepcion)."""
    t = _sin_tildes(_texto(v).upper())
    if not t:
        return None, "sin centro asignado"
    if "PUENTESAUCO" in t:
        return None, "centro mixto con Puentesaúco — confirmar censo con Dirección"
    if "FUENTECILLAS" not in t:
        return None, f"centro no reconocido: {t[:40]}"
    return ("CD" if "DIA" in t else "RESIDENCIA"), None


def normalizar_grado_dep(v: str) -> str:
    t = _texto(v).upper().replace("GRADO", "").strip()
    return {"1": "I", "I": "I", "2": "II", "II": "II", "3": "III", "III": "III"}.get(t, "no_valorado")


def normalizar_forma_com(v: str) -> str:
    t = _sin_tildes(_texto(v).lower())
    tiene_tel = "tel" in t
    tiene_mail = "mail" in t or "@" in t or "email" in t
    if tiene_tel and tiene_mail:
        return "Teléfono y email"
    if tiene_mail:
        return "Email"
    if tiene_tel:
        return "Teléfono"
    return ""


def normalizar_relacion(v: str) -> str:
    t = _sin_tildes(_texto(v).upper())
    for clave, slug in RELACIONES.items():
        if clave in t:
            return slug
    return "otro"


def trocear_lista(v: str) -> list[str]:
    """Texto libre multivalor → lista de entradas ≥ 4 caracteres."""
    bruto = str(v or "")
    partes = re.split(r"[\n;·]+|(?<=[a-záéíóú)])\.\s+(?=[A-ZÁÉÍÓÚ])", bruto)
    return [re.sub(r"\s+", " ", p).strip(" .") for p in partes
            if p and len(p.strip(" .")) >= 4]


def procesar(ruta: Path):
    """→ (registros, excepciones, avisos_globales). Sin datos por consola."""
    from openpyxl import load_workbook  # import tardío: la capa ORM no lo necesita

    wb = load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = list(ws.iter_rows(values_only=True))
    registros, excepciones = [], []

    for n_fila, fila in enumerate(filas[1:], 2):
        if not any(c is not None and str(c).strip() for c in fila):
            continue
        def col(campo):
            i = COLS[campo]
            return fila[i] if len(fila) > i else None

        etiqueta = f"{_texto(col('apellidos'))}, {_texto(col('nombre'))}"
        problemas_bloqueantes, avisos = [], []

        apellido_1, apellido_2 = separar_apellidos(_texto(col("apellidos")))
        if not apellido_1 or not _texto(col("nombre")):
            problemas_bloqueantes.append("nombre o apellidos vacíos")

        f_nac = normalizar_fecha(col("f_nac"))
        if not f_nac:
            problemas_bloqueantes.append("fecha de nacimiento inválida")

        f_alta = normalizar_fecha(col("f_incorporacion"))
        if not f_alta:
            avisos.append("fecha de incorporación inválida — se usará 01/01/2000 provisional")
            f_alta = date(2000, 1, 1)

        centro, motivo_centro = normalizar_centro(col("centro"))
        if motivo_centro:
            problemas_bloqueantes.append(motivo_centro)

        dni = _texto(col("dni")).upper().replace(" ", "").replace("-", "")
        if dni and not validar_nif(dni):
            avisos.append(f"NIF con letra que no cuadra: revisar")

        pct = None
        try:
            pct_txt = _texto(col("pct_disc")).replace("%", "").replace(",", ".")
            if pct_txt:
                pct = int(float(pct_txt))
        except ValueError:
            avisos.append("% discapacidad no numérico — se omite")

        if problemas_bloqueantes:
            excepciones.append({"fila": n_fila, "persona": etiqueta,
                                "motivos": "; ".join(problemas_bloqueantes)})
            continue

        notas = []
        if _texto(col("medicacion")):
            fecha_med = _texto(col("fecha_cambio_med"))
            notas.append(f"MEDICACIÓN HABITUAL (migrado{' · últ. cambio ' + fecha_med if fecha_med else ''}): {_texto(col('medicacion'))}")
        if _texto(col("quirurgicas")):
            notas.append(f"INTERVENCIONES QUIRÚRGICAS (migrado): {_texto(col('quirurgicas'))}")
        if _texto(col("icap")):
            notas.append(f"ICAP (migrado): {_texto(col('icap'))}{' · ' + _texto(col('f_icap')) if _texto(col('f_icap')) else ''}")
        if _texto(col("expediente")):
            notas.append(f"REFERENCIA/EXPEDIENTE (origen): {_texto(col('expediente'))}")

        registros.append({
            "fila": n_fila,
            "persona": {
                "nombre": _texto(col("nombre")), "apellido_1": apellido_1,
                "apellido_2": apellido_2, "dni_nie": dni,
                "fecha_nacimiento": f_nac, "fecha_alta": f_alta,
                "direccion_calle": _texto(col("direccion")),
                "direccion_cp": _texto(col("cp")),
                "direccion_municipio": _texto(col("ciudad")),
                "direccion_provincia": _texto(col("provincia")),
                "notas_relevantes": "\n".join(notas),
            },
            "centro": centro,
            "sanitario": {"tsi": _texto(col("tsi")), "nuss": _texto(col("nuss")),
                          "tsi_caducidad": normalizar_fecha(col("cad_dni"))},
            "perfil_mayor": {
                "grado_dependencia": normalizar_grado_dep(col("grado_dep")),
                "fecha_valoracion_bvd": normalizar_fecha(col("f_res_dep")),
                "inicio_at": normalizar_fecha(col("f_res_dep")) or f_alta,
            },
            "perfil_di": ({"tipo_discapacidad": "no_valorada",
                           "grado_discapacidad_pct": pct, "inicio_at": f_alta}
                          if pct is not None else None),
            "enfermedades": trocear_lista(col("enf_cronicas")),
            "diagnosticos": [d for d in (_texto(col("diag_1")), _texto(col("diag_2"))) if d],
            "f_validez_diag": normalizar_fecha(col("f_validez_diag")),
            "contacto": ({
                "nombre": _texto(col("contacto_nombre")),
                "relacion": normalizar_relacion(col("parentesco")),
                "relacion_texto": _texto(col("parentesco")),
                "movil": _texto(col("tel1")), "telefono_fijo": _texto(col("tel2")),
                "email": _texto(col("mail")),
                "forma_comunicacion": normalizar_forma_com(col("forma_com")),
            } if _texto(col("contacto_nombre")) else None),
            "avisos": avisos,
        })

    avisos_globales = [
        "«Ingresos mensuales» NO importado (minimización art. 5 RGPD).",
        "«Prestaciones dependencia» ignorado (columna vacía en origen).",
        "Teléfonos y mail asignados al CONTACTO principal (convención Ixis) — validar.",
    ]
    return registros, excepciones, avisos_globales


def generar_informe(registros, excepciones, avisos_globales, salida: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.column_dimensions["A"].width = 90
    lineas = [f"Informe de importación Fuentecillas — {datetime.now():%d/%m/%Y %H:%M}",
              f"Personas importables: {len(registros)}",
              f"Excepciones (no se importan aún): {len(excepciones)}",
              f"Con avisos no bloqueantes: {sum(1 for r in registros if r['avisos'])}",
              ""] + ["· " + a for a in avisos_globales]
    for i, l in enumerate(lineas, 1):
        c = ws.cell(row=i, column=1, value=l)
        if i == 1:
            c.font = Font(bold=True, size=13, color="047857")

    ws2 = wb.create_sheet("Detalle")
    ws2.append(["Fila", "Persona", "Centro", "Grado dep.", "% Disc", "Enfermedades (nº)",
                "Contacto", "Avisos"])
    for col in range(1, 9):
        ws2.cell(row=1, column=col).font = Font(bold=True)
    for r in registros:
        p = r["persona"]
        ws2.append([r["fila"], f"{p['apellido_1']} {p['apellido_2']}, {p['nombre']}".strip(),
                    r["centro"], r["perfil_mayor"]["grado_dependencia"],
                    (r["perfil_di"] or {}).get("grado_discapacidad_pct"),
                    len(r["enfermedades"]),
                    (r["contacto"] or {}).get("nombre", "—"),
                    " | ".join(r["avisos"])])
        if r["avisos"]:
            for col in range(1, 9):
                ws2.cell(row=ws2.max_row, column=col).fill = PatternFill("solid", fgColor="FEF3C7")

    ws3 = wb.create_sheet("Excepciones")
    ws3.append(["Fila", "Persona", "Motivos"])
    for e in excepciones:
        ws3.append([e["fila"], e["persona"], e["motivos"]])

    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    regs, excs, avisos = procesar(Path(sys.argv[1]))
    generar_informe(regs, excs, avisos, Path(sys.argv[2]))
    print(f"Importables: {len(regs)} · Excepciones: {len(excs)} · "
          f"Con avisos: {sum(1 for r in regs if r['avisos'])}")
    print(f"Informe: {sys.argv[2]}")
